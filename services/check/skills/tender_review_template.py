from __future__ import annotations

import io
import base64
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

TEMPLATE = [
    {"name": "项目信息", "tab": "1F4E78", "title": "项目信息表", "subtitle": "投标文件审查固定模板", "headers": ["项目", "内容"], "data_key": None},
    {"name": "审查概要与风险", "tab": "C00000", "title": "审查概要与风险总览", "subtitle": "高风险红 / 中风险橙 / 低风险绿", "headers": ["序号", "风险类别", "风险点描述", "风险等级", "涉及条款", "影响后果", "应对建议"], "data_key": "overview"},
    {"name": "评分项分析与预测", "tab": "00B050", "title": "评分项分析与预测得分", "subtitle": "每项必须给出单一预测得分", "headers": ["序号", "评审因素", "满分", "预测得分", "预计失分", "预测理由", "响应状态", "响应依据", "差距分析", "补充建议"], "data_key": "scoring"},
    {"name": "废标项核查", "tab": "ED7D31", "title": "废标项核查", "subtitle": "投标递交前必须逐项处置", "headers": ["序号", "检查项", "招标要求", "投标响应", "核查结果", "风险等级", "修改建议"], "data_key": "disqualification_summary"},
    {"name": "分项报价", "tab": "92D050", "title": "分项报价核对", "subtitle": "报价空白和超限价优先核查", "headers": ["序号", "服务名称", "数量", "单价", "小计", "核对结果"], "data_key": "pricing"},
    {"name": "交付时间对比", "tab": "0070C0", "title": "交付时间对比", "subtitle": "招标要求与投标承诺逐项对比", "headers": ["序号", "事项", "招标要求", "投标承诺", "差异", "风险等级"], "data_key": "delivery"},
    {"name": "时间节点", "tab": "7030A0", "title": "关键时间节点", "subtitle": "投标与履约关键节点", "headers": ["序号", "时间节点", "日期/时间", "状态", "备注"], "data_key": "milestones"},
    {"name": "废标项核对(BidMaster)", "tab": "FF6600", "title": "废标项核对(BidMaster)", "subtitle": "汇总商务线与技术线废标项", "headers": ["序号", "条款类型", "条款内容", "原文出处", "命中关键句", "投标书是否响应", "响应内容", "风险等级", "修改建议"], "data_key": "disqualification"},
    {"name": "▲参数核对", "tab": "808000", "title": "▲参数核对", "subtitle": "标识参数逐条回原文核对", "headers": ["序号", "标识", "参数描述", "类别", "投标书响应状态", "响应内容", "偏离情况", "建议"], "data_key": "star_params"},
    {"name": "证明材料清单", "tab": "2E75B6", "title": "证明材料清单", "subtitle": "证明材料、页码和状态逐项核对", "headers": ["序号", "材料名称", "要求描述", "原文出处", "命中关键句", "投标书是否包含", "页码/章节", "备注"], "data_key": "materials"},
    {"name": "时间节点核对", "tab": "00B0F0", "title": "时间节点核对", "subtitle": "招标要求与投标响应逐项核对", "headers": ["序号", "节点类型", "时间/期限", "要求描述", "投标书是否响应", "响应内容", "风险提示"], "data_key": "timeline"},
    {"name": "合同条款要点", "tab": "7F7F7F", "title": "合同条款要点", "subtitle": "中标后约束和履约风险", "headers": ["序号", "条款主题", "约束要点", "限制性原话", "原文出处", "投标书响应情况", "中标后风险", "应对建议"], "data_key": "contract_terms"},
]


def build_excel(
    data: dict[str, list[dict]],
    dimension_headers: dict[str, list[str]],
    item_key: object,
    company: str,
    school: str,
) -> str:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in TEMPLATE:
        worksheet = workbook.create_sheet(spec["name"])
        worksheet.sheet_properties.tabColor = spec["tab"]
        worksheet.sheet_view.showGridLines = False
        write_header(worksheet, spec, data, company, school)
        rows = rows_for(spec, data, dimension_headers, item_key)
        for row_number, values in enumerate(rows, 5):
            for column_number, value in enumerate(values, 1):
                cell = worksheet.cell(row_number, column_number, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = THIN
            risk = risk_level(values)
            if risk:
                worksheet.cell(row_number, 1).fill = PatternFill(
                    "solid",
                    fgColor={"high": "F4CCCC", "medium": "FCE5CD", "low": "D9EAD3"}.get(risk, "FFFFFF"),
                )
        format_sheet(worksheet, len(spec["headers"]), len(rows) + 4)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def write_header(worksheet, spec, data, company, school) -> None:
    headers = spec["headers"]
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    worksheet.cell(1, 1, spec["title"])
    worksheet.cell(2, 1, subtitle(spec, data, company, school))
    for row_number in (1, 2):
        for column_number in range(1, len(headers) + 1):
            cell = worksheet.cell(row_number, column_number)
            if row_number == 1:
                cell.fill = PatternFill("solid", fgColor=spec["tab"])
                cell.font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
            else:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=spec["tab"])
            cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 22
    worksheet.row_dimensions[3].height = 8
    if spec["name"] == "项目信息":
        return
    for column_number, header in enumerate(headers, 1):
        cell = worksheet.cell(4, column_number, header)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    worksheet.row_dimensions[4].height = 28
    worksheet.freeze_panes = "A5"


def subtitle(spec, data, company, school) -> str:
    if spec["name"] == "项目信息":
        return f"{company} · {school} · 生成时间 {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    if spec["name"] == "评分项分析与预测":
        scores = [numeric_score(item.get("predicted_score")) for item in data.get("scoring", []) if isinstance(item, dict)]
        scores = [score for score in scores if score is not None]
        if scores:
            return f"单项预测合计：{sum(scores):g} 分 · 每项为单一预测值"
        return "每项必须给出单一预测得分"
    return spec["subtitle"]


def rows_for(spec, data, dimension_headers, item_key) -> list[list]:
    name = spec["name"]
    if name == "项目信息":
        return project_info_rows(data)
    if name == "审查概要与风险":
        return overview_rows(spec, data)
    if name == "废标项核查":
        return disqualification_checklist_rows(spec, data)
    if name == "分项报价":
        return placeholder_rows(spec, data.get("pricing", []), "AI未提取到结构化分项报价")
    if name == "交付时间对比":
        return placeholder_rows(spec, data.get("delivery", []), "AI未提取到结构化交付时间")
    if name == "时间节点":
        return milestone_rows(spec, data)
    if name == "废标项核对(BidMaster)":
        return dimension_rows("disqualification", spec["headers"], data, item_key)
    return dimension_rows(spec["data_key"], spec["headers"], data, item_key)


def project_info_rows(data) -> list[list]:
    info = data.get("project_info", [])
    rows = []
    seen = set()
    for item in info:
        if not isinstance(item, dict):
            continue
        label = str(item.get("label", item.get("项目", ""))).strip()
        value = str(item.get("value", item.get("内容", ""))).strip()
        if not label or (label, value) in seen:
            continue
        seen.add((label, value))
        rows.append([label, value])
    return rows


def overview_rows(spec, data) -> list[list]:
    rows = []
    dimensions = {
        "disqualification": "废标项核对",
        "scoring": "评分项分析与预测",
        "star_params": "▲参数核对",
        "materials": "证明材料清单",
        "timeline": "时间节点核对",
        "contract_terms": "合同条款要点",
        "pricing": "分项报价",
        "delivery": "交付时间对比",
    }
    total_items = 0
    high_items = 0
    review_items = 0
    for key, title in dimensions.items():
        items = data.get(key, [])
        dimension_items = [item for item in items if isinstance(item, dict)]
        total_items += len(dimension_items)
        high = sum(1 for item in items if isinstance(item, dict) and item.get("risk_level") == "high")
        high_items += high
        risk = "高" if high else "低"
        if key == "disqualification":
            unresolved = sum(1 for item in dimension_items if item.get("response_found") is False or item.get("risk_level") == "high")
            pending = sum(1 for item in dimension_items if item.get("response_found") is not True)
            risk = "高" if unresolved else ("中" if pending else "低")
        elif key == "star_params":
            negative = sum(1 for item in dimension_items if item.get("response_status") in {"negative_deviation", "missing"})
            vague = sum(1 for item in dimension_items if item.get("response_status") == "vague")
            risk = "高" if negative else ("中" if vague else "低")
        elif key == "materials":
            missing = sum(1 for item in dimension_items if item.get("included") is False)
            pending = sum(1 for item in dimension_items if item.get("included") is not True)
            risk = "高" if missing else ("中" if pending else "低")
        elif key == "timeline":
            missing = sum(1 for item in dimension_items if item.get("responded") is False)
            pending = sum(1 for item in dimension_items if item.get("responded") is not True)
            risk = "高" if missing else ("中" if pending else "低")
        elif key == "pricing":
            bad = sum(1 for item in dimension_items if any(word in str(item.get("核对结果", "")) for word in ("缺失", "不一致")))
            pending = sum(1 for item in dimension_items if "待人工复核" in str(item.get("核对结果", "")))
            risk = "高" if bad else ("中" if pending else "低")
        elif key == "delivery":
            delayed = sum(1 for item in dimension_items if any(word in str(item.get("差异", "")) for word in ("延迟", "未承诺")))
            pending = sum(1 for item in dimension_items if any(word in str(item.get("差异", "")) for word in ("模糊", "待确认")))
            risk = "高" if delayed else ("中" if pending else "低")
        if key == "contract_terms":
            finding = "重点管理中标后合同履约风险"
        else:
            finding = {
                "disqualification": "存在未响应或高风险废标项，必须逐条处置" if risk != "低" else "未发现高风险废标项",
                "scoring": "以预测得分和差距分析为准，优先补强低分项",
                "star_params": "优先处理负偏离和模糊响应" if risk != "低" else "星号参数响应良好",
                "materials": "存在缺失或待确认材料，需补齐证明文件" if risk != "低" else "证明材料响应完整",
                "timeline": "存在未响应时间节点，需逐项确认" if risk != "低" else "时间节点响应完整",
                "pricing": "存在缺失、不一致或待复核报价，需重新核对" if risk != "低" else "分项报价无明显异常",
                "delivery": "存在未承诺、延迟或模糊交付时间，需修正承诺" if risk != "低" else "交付时间承诺无重大差异",
            }[key]
        rows.append([
            len(rows) + 1,
            title,
            f"提取 {len(dimension_items)} 条" + (f"，高风险 {high} 条" if high else ""),
            risk,
            "AI 全量审查结果",
            finding,
            "优先处理高风险项" if risk == "高" else ("递交前人工复核待确认项" if risk == "中" else "保持现有响应材料"),
        ])
        if key != "contract_terms":
            review_items += len(dimension_items)
    data["_summary_stats"] = [{
        "审查维度数": str(len(rows)),
        "审查条目数": str(review_items),
        "高风险条目数": str(high_items),
        "总提取条目数": str(total_items),
    }]
    return rows or [[1, "审查结果", "本轮未提取到结构化风险", "低", "全量审查", "暂无", "仍建议递交前人工复核"]]


def disqualification_checklist_rows(spec, data) -> list[list]:
    rows = []
    for item in data.get("disqualification", []):
        if not isinstance(item, dict):
            continue
        response_found = item.get("response_found")
        status = "❌ 废标风险" if item.get("risk_level") == "high" else "✅ 合格" if response_found is True else "⚠️ 待确认"
        rows.append([
            item.get("seq", len(rows) + 1),
            item.get("clause_type", "废标项"),
            item.get("clause_content", ""),
            item.get("response_content", "未找到"),
            status,
            item.get("risk_level", ""),
            item.get("suggestion", ""),
        ])
    return rows or [[1, "废标项核查", "未提取到结构化废标项", "待人工复核", "⚠️ 待确认", "中", "递交前按招标文件逐项复核"]]


def milestone_rows(spec, data) -> list[list]:
    rows = []
    for item in data.get("timeline", []):
        if not isinstance(item, dict):
            continue
        responded = item.get("responded")
        status = "已响应" if responded is True else "未响应" if responded is False else "待确认"
        rows.append([item.get("seq", len(rows) + 1), item.get("node_type", ""), item.get("time_value", ""), status, item.get("risk_note", "")])
    return rows or [[1, "关键时间节点", "待补充", "待确认", "AI未提取到结构化时间节点"]]


def dimension_rows(key, headers, data, item_key) -> list[list]:
    result = []
    for item in data.get(key, []):
        if not isinstance(item, dict):
            continue
        row = []
        for column_number, header in enumerate(headers, 1):
            field = item_key(key, column_number)
            value = item.get(field, item.get(header, ""))
            if isinstance(value, bool):
                value = "是" if value else "否"
            if key == "scoring" and header == "预计失分":
                score = numeric_score(item.get("score"))
                predicted = numeric_score(item.get("predicted_score"))
                value = max(score - predicted, 0) if score is not None and predicted is not None else ""
            if key == "scoring" and header == "响应状态":
                value = {"answered": "已响应", "partial": "部分响应", "missing": "未响应"}.get(value, value or "待确认")
            if key == "star_params" and header == "响应状态":
                value = {"compliant": "已响应", "negative_deviation": "负偏离", "vague": "模糊响应", "missing": "未响应"}.get(value, value or "待确认")
            if key in {"disqualification", "materials", "timeline"} and header in {"投标书是否响应", "投标书是否包含"}:
                value = "是" if value is True else "否" if value is False else "待确认"
            if key == "delivery" and header == "风险等级":
                value = {"high": "高", "medium": "中", "low": "低"}.get(value, value or "待确认")
            if key == "scoring" and header == "预测得分" and value == "":
                value = "待补充单一预测得分"
            row.append("" if value is None else value)
        result.append(row)
    return result


def placeholder_rows(spec, items, fallback_text="AI未提取到结构化数据") -> list[list]:
    result = []
    for item in items or []:
        if not isinstance(item, dict):
            continue
        row = []
        for header in spec["headers"]:
            value = item.get(header, "")
            if isinstance(value, bool):
                value = "是" if value else "否"
            row.append("" if value is None else value)
        result.append(row)
    return result or [[1] + [fallback_text] * (len(spec["headers"]) - 1)]


def risk_level(row) -> str:
    text = " ".join(str(value) for value in row)
    lowered = text.lower()
    if any(token in text for token in ("高", "❌", "废标风险", "未响应", "未覆盖", "未提供")) or "high" in lowered:
        return "high"
    if any(token in text for token in ("中", "⚠", "待确认", "部分响应", "需人工复核")) or "medium" in lowered or "partial" in lowered:
        return "medium"
    if any(token in text for token in ("低", "✅", "合格", "已响应", "已覆盖", "已提供")) or "low" in lowered or "answered" in lowered:
        return "low"
    return ""


def numeric_score(value):
    try:
        return float(str(value).strip().removesuffix("分"))
    except (TypeError, ValueError, AttributeError):
        return None


def format_sheet(worksheet, column_count, row_count) -> None:
    for column_number in range(1, column_count + 1):
        lengths = []
        for row_number in range(4, row_count + 1):
            value = str(worksheet.cell(row_number, column_number).value or "")
            lengths.append(sum(2 if ord(char) > 127 else 1 for char in value[:120]))
        worksheet.column_dimensions[get_column_letter(column_number)].width = min(max(max(lengths, default=0) + 4, 12), 48)
    worksheet.auto_filter.ref = f"A4:{get_column_letter(column_count)}{max(row_count, 4)}"
