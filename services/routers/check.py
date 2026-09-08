from __future__ import annotations

import asyncio
import importlib
import functools
import logging
import re
import tempfile
import os

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from services.database import get_db
from services.models import (
    Project, Document, Analysis, Chapter, CheckReport,
    ProjectStatus, CheckType,
)
from services.llm_factory import get_agent_gateway
from core.skill_engine.base import SkillContext
from core.task_manager import TaskManager
from core.settings import get_settings

logger = logging.getLogger(__name__)
router = APIRouter()

MAX_UPLOAD_BYTES = 50 * 1024 * 1024  # 50MB


def _truncate_text(text: str, max_chars: int) -> str:
    if not text or max_chars <= 0 or len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[... 内容已截断，已省略后续部分 ...]"


def _check_max_concurrent() -> int:
    """检查链的 LLM 并发度。默认 8，可用 BMP_CHECK_MAX_CONCURRENT 覆盖。

    全面检查一次要跑 15 个 skill，裸 asyncio.gather 会把 15 路请求同时压到网关上；
    一旦触发限流就是 15 项一起失败，前端表现为「全是异常」。限流后墙钟时间略增，
    但换来的是可预期的成功率——与解读链 BMP_INTERPRET_MAX_CONCURRENT 同一套做法。
    """
    try:
        value = int(os.getenv("BMP_CHECK_MAX_CONCURRENT", "8"))
    except ValueError:
        return 8
    return max(1, min(value, 15))


async def _get_tender_and_bid_text(project_id: str, db: AsyncSession):
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    max_chars = get_settings().tender_text_max_chars

    tender_text = ""
    if project.tender_doc_id:
        doc_result = await db.execute(
            select(Document).where(Document.id == project.tender_doc_id)
        )
        doc = doc_result.scalar_one_or_none()
        if doc and doc.parsed_content:
            tender_text = _truncate_text(doc.parsed_content, max_chars)

    bid_docs = await db.execute(
        select(Document).where(
            Document.project_id == project.id,
            Document.type == "bid",
        )
    )
    bid_doc = bid_docs.scalars().first()
    bid_text = bid_doc.parsed_content if bid_doc and bid_doc.parsed_content else ""

    if not bid_text:
        chapter_result = await db.execute(
            select(Chapter).where(Chapter.project_id == project.id)
        )
        chapters = chapter_result.scalars().all()
        if chapters:
            bid_text = "\n\n".join(
                f"## {ch.title}\n{ch.content or ''}" for ch in chapters if ch.content
            )

    bid_text = _truncate_text(bid_text, max_chars)

    return project, tender_text, bid_text


@router.post("/{project_id}/compliance")
async def check_compliance(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

    if not tender_text or not bid_text:
        raise HTTPException(status_code=400, detail="招标文件或投标文件内容为空")

    from services.check.skills.compliance_check_skill import ComplianceCheckSkill

    gateway = await get_agent_gateway(db, "check")
    skill = ComplianceCheckSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={"tender_text": tender_text, "bid_text": bid_text},
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.COMPLIANCE,
            results=skill_result.data,
            risk_level="high" if skill_result.data.get("has_critical_issues") else "low",
            summary={
                "total": skill_result.data.get("total_requirements", 0),
                "compliant": skill_result.data.get("compliant", 0),
                "non_compliant": skill_result.data.get("non_compliant", 0),
            },
        )
        db.add(report)
        await db.flush()

    return {
        "success": skill_result.success,
        "data": skill_result.data,
        "error": skill_result.error,
        "warnings": skill_result.warnings,
    }


@router.post("/{project_id}/disqualification")
async def check_disqualification(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

    if not tender_text or not bid_text:
        raise HTTPException(status_code=400, detail="招标文件或投标文件内容为空")

    analysis_result = await db.execute(
        select(Analysis).where(Analysis.project_id == project.id)
    )
    analysis = analysis_result.scalar_one_or_none()
    dq_clauses = []
    if analysis and analysis.dimensions:
        dq_data = analysis.dimensions.get("disqualification", {})
        if dq_data and not dq_data.get("error"):
            dq_clauses = dq_data.get("clauses", [])

    from services.check.skills.disqualification_check_skill import DisqualificationCheckSkill

    gateway = await get_agent_gateway(db, "check")
    skill = DisqualificationCheckSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={
            "tender_text": tender_text,
            "bid_text": bid_text,
            "disqualification_clauses": dq_clauses,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.DISQUALIFICATION,
            results=skill_result.data,
            risk_level="high" if skill_result.data.get("missing", 0) > 0 else "low",
        )
        db.add(report)
        await db.flush()

    return {
        "success": skill_result.success,
        "data": skill_result.data,
        "error": skill_result.error,
        "warnings": skill_result.warnings,
    }


@router.post("/{project_id}/qualification")
async def check_qualification(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

    from services.check.skills.qualification_check_skill import QualificationCheckSkill

    analysis_result = await db.execute(
        select(Analysis).where(Analysis.project_id == project.id)
    )
    analysis = analysis_result.scalar_one_or_none()
    timeline = analysis.dimensions.get("timeline", {}) if analysis and analysis.dimensions else {}
    bid_deadline = timeline.get("投标截止日", "") if isinstance(timeline, dict) else ""

    gateway = await get_agent_gateway(db, "check")
    skill = QualificationCheckSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={
            "tender_text": tender_text,
            "bid_text": bid_text,
            "bid_deadline": bid_deadline,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.QUALIFICATION,
            results=skill_result.data,
            risk_level=skill_result.data.get("risk_level", "low"),
        )
        db.add(report)
        await db.flush()

    return {
        "success": skill_result.success,
        "data": skill_result.data,
        "error": skill_result.error,
        "warnings": skill_result.warnings,
    }


@router.post("/{project_id}/pricing")
async def check_pricing(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

    analysis_result = await db.execute(
        select(Analysis).where(Analysis.project_id == project.id)
    )
    analysis = analysis_result.scalar_one_or_none()
    max_price = None
    if analysis and analysis.dimensions:
        project_info = analysis.dimensions.get("project_info", {})
        if isinstance(project_info, dict):
            budget_str = project_info.get("预算金额", "")
            if budget_str:
                match = re.search(r"[\d.]+", str(budget_str))
                if match:
                    max_price = float(match.group())

    from services.check.skills.pricing_check_skill import PricingCheckSkill

    gateway = await get_agent_gateway(db, "check")
    skill = PricingCheckSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={
            "tender_text": tender_text,
            "bid_text": bid_text,
            "max_price": max_price,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.PRICING,
            results=skill_result.data,
            risk_level=skill_result.data.get("risk_level", "low"),
        )
        db.add(report)
        await db.flush()

    return {
        "success": skill_result.success,
        "data": skill_result.data,
        "error": skill_result.error,
        "warnings": skill_result.warnings,
    }


@router.post("/{project_id}/fit-score")
async def check_fit_score(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

    if not tender_text or not bid_text:
        raise HTTPException(status_code=400, detail="招标文件或投标文件内容为空")

    from services.check.skills.fit_score_skill import FitScoreSkill

    gateway = await get_agent_gateway(db, "check")
    skill = FitScoreSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={"tender_text": tender_text, "bid_text": bid_text},
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.FIT_SCORE,
            results=skill_result.data,
        )
        db.add(report)
        await db.flush()

    return {
        "success": skill_result.success,
        "data": skill_result.data,
        "error": skill_result.error,
        "warnings": skill_result.warnings,
    }


@router.post("/{project_id}/selfcheck")
async def run_selfcheck(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

    check_results = {}

    if tender_text and bid_text:
        gateway = await get_agent_gateway(db, "check")

        # Extract max_price for pricing check
        analysis_result = await db.execute(
            select(Analysis).where(Analysis.project_id == project.id)
        )
        analysis = analysis_result.scalar_one_or_none()
        max_price = None
        if analysis and analysis.dimensions:
            project_info = analysis.dimensions.get("project_info", {})
            if isinstance(project_info, dict):
                budget_str = project_info.get("预算金额", "")
                if budget_str:
                    match = re.search(r"[\d.]+", str(budget_str))
                    if match:
                        max_price = float(match.group())

        # Extract bid_deadline for qualification check
        timeline = analysis.dimensions.get("timeline", {}) if analysis and analysis.dimensions else {}
        bid_deadline = timeline.get("投标截止日", "") if isinstance(timeline, dict) else ""

        from services.check.skills.compliance_check_skill import ComplianceCheckSkill
        compliance_skill = ComplianceCheckSkill()
        compliance_ctx = SkillContext(
            project_id=project_id, db=db, llm=gateway,
            parameters={"tender_text": tender_text, "bid_text": bid_text},
        )
        compliance_result = await compliance_skill.safe_execute(compliance_ctx)
        if compliance_result.success:
            check_results["compliance_check"] = compliance_result.data

        from services.check.skills.pricing_check_skill import PricingCheckSkill
        pricing_skill = PricingCheckSkill()
        pricing_ctx = SkillContext(
            project_id=project_id, db=db, llm=gateway,
            parameters={"tender_text": tender_text, "bid_text": bid_text, "max_price": max_price},
        )
        pricing_result = await pricing_skill.safe_execute(pricing_ctx)
        if pricing_result.success:
            check_results["pricing_check"] = pricing_result.data

        from services.check.skills.qualification_check_skill import QualificationCheckSkill
        qual_skill = QualificationCheckSkill()
        qual_ctx = SkillContext(
            project_id=project_id, db=db, llm=gateway,
            parameters={"tender_text": tender_text, "bid_text": bid_text, "bid_deadline": bid_deadline},
        )
        qual_result = await qual_skill.safe_execute(qual_ctx)
        if qual_result.success:
            check_results["qualification_check"] = qual_result.data

    from services.check.skills.selfcheck_list_skill import SelfcheckListSkill

    skill = SelfcheckListSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=await get_agent_gateway(db, "check"),
        parameters={"check_results": check_results},
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.SELFCHECK,
            results=skill_result.data,
            risk_level="low" if skill_result.data.get("can_submit") else "high",
        )
        db.add(report)
        await db.flush()

    return {
        "success": skill_result.success,
        "data": skill_result.data,
        "error": skill_result.error,
        "warnings": skill_result.warnings,
    }


@router.post("/{project_id}/full-check")
async def full_check(project_id: str, db: AsyncSession = Depends(get_db)):
    """Submit full check as async task. Returns task_id for polling."""
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    project.status = ProjectStatus.CHECKING
    await db.flush()
    await db.commit()

    tm = TaskManager.instance()
    task = await tm.submit("full_check", _do_full_check, project_id)

    return {
        "task_id": task.task_id,
        "status": "pending",
        "message": "全面检查任务已提交，请通过 GET /check/task/{task_id} 查询进度",
    }


async def _do_full_check(project_id: str):
    """Background worker: run 15 check skills in parallel and persist results."""
    from services.database import async_session

    session_factory = async_session()
    async with session_factory() as db:
        try:
            project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

            if not tender_text or not bid_text:
                raise ValueError("招标文件或投标文件内容为空")

            gateway = await get_agent_gateway(db, "check")

            # Extract supplementary parameters
            analysis_result = await db.execute(
                select(Analysis).where(Analysis.project_id == project.id)
            )
            analysis = analysis_result.scalar_one_or_none()
            max_price = None
            bid_deadline = ""
            if analysis and analysis.dimensions:
                project_info = analysis.dimensions.get("project_info", {})
                if isinstance(project_info, dict):
                    budget_str = project_info.get("预算金额", "")
                    if budget_str:
                        match = re.search(r"[\d.]+", str(budget_str))
                        if match:
                            max_price = float(match.group())
                timeline = analysis.dimensions.get("timeline", {})
                if isinstance(timeline, dict):
                    bid_deadline = timeline.get("投标截止日", "")

            # 并发执行但必须限流，理由见 _check_max_concurrent()。
            semaphore = asyncio.Semaphore(_check_max_concurrent())

            async def _run_skill(ct: str, params: dict) -> tuple[str, dict]:
                async with semaphore:
                    return await _exec_check_skill(ct, params, gateway, project_id)

            # Build project_facts for consistency check (same as standalone endpoint)
            project_facts = {}
            if analysis and analysis.dimensions:
                project_facts = {
                    "project_name": project.name,
                    "dimensions": analysis.dimensions,
                }

            # 15 项检查及各自的补充参数（项目模式能从 Analysis 拿到上下文）
            skill_tasks = [
                (
                    ct,
                    _full_check_params(
                        ct,
                        tender_text,
                        bid_text,
                        bid_deadline=bid_deadline,
                        max_price=max_price,
                        project_facts=project_facts,
                    ),
                )
                for ct in _FULL_CHECK_TYPES
            ]

            # Run all skills in parallel
            results = await asyncio.gather(
                *[_run_skill(ct, params) for ct, params in skill_tasks],
                return_exceptions=True,
            )

            all_results = {}
            for r in results:
                if isinstance(r, Exception):
                    logger.warning(f"[full_check] skill 异常被 asyncio.gather 捕获: {r}")
                    continue
                ct, result_dict = r
                all_results[ct] = result_dict

            has_critical = any(
                r.get("data", {}).get("has_critical_issues") or r.get("data", {}).get("risk_level") == "high"
                for r in all_results.values()
                if r.get("success") and isinstance(r.get("data"), dict)
            )

            # Persist report
            report = CheckReport(
                project_id=project.id,
                type=CheckType.FULL_CHECK,
                results=all_results,
                risk_level="high" if has_critical else "low",
                summary={"checks_run": len(all_results), "has_critical": has_critical},
            )
            db.add(report)
            project.status = ProjectStatus.COMPLETED
            await db.flush()
            await db.commit()

            return {
                "success": True,
                "data": all_results,
                "has_critical": has_critical,
            }
        except Exception as e:
            logger.error(f"[full_check] 异常: {e}")
            try:
                project_result = await db.execute(
                    select(Project).where(Project.id == project_id)
                )
                proj = project_result.scalar_one_or_none()
                if proj:
                    proj.status = ProjectStatus.COMPLETED
                    await db.commit()
            except Exception:
                pass
            raise  # Let TaskManager correctly mark the task as FAILED


@router.get("/task/{task_id}")
async def get_check_task_status(task_id: str):
    """Poll endpoint for async check tasks."""
    tm = TaskManager.instance()
    task = tm.get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    return task.to_dict()


@router.get("/{project_id}/reports")
async def list_check_reports(project_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    reports_result = await db.execute(
        select(CheckReport).where(CheckReport.project_id == project.id)
    )
    reports = reports_result.scalars().all()

    return {"reports": [
        {
            "id": str(r.id),
            "type": r.type.value if isinstance(r.type, CheckType) else r.type,
            "risk_level": r.risk_level,
            "summary": r.summary,
            "created_at": r.created_at.isoformat() if r.created_at else None,
        }
        for r in reports
    ]}


@router.post("/{project_id}/deposit")
async def check_deposit(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.deposit_check_skill import DepositCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = DepositCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.DEPOSIT, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


async def _parse_uploaded_file(file: UploadFile) -> str:
    suffix = os.path.splitext(file.filename or "")[1].lower()
    content_bytes = await file.read()
    if len(content_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限",
        )

    if suffix == ".txt" or suffix == ".md":
        return content_bytes.decode("utf-8", errors="replace")

    if suffix == ".docx":
        tmp_path = ""
        try:
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
                tmp.write(content_bytes)
                tmp_path = tmp.name

            from docx import Document as DocxDocument
            doc = DocxDocument(tmp_path)
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    style_name = (para.style.name or "").lower() if para.style else ""
                    if "heading" in style_name or "标题" in style_name:
                        level = "1"
                        for ch in style_name:
                            if ch.isdigit():
                                level = ch
                                break
                        paragraphs.append(f"{'#' * int(level)} {text}")
                    else:
                        paragraphs.append(text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        paragraphs.append(" | ".join(cells))

            return "\n\n".join(paragraphs)
        except Exception as e:
            return content_bytes.decode("utf-8", errors="replace")
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)

    if suffix == ".pdf":
        tmp_path = ""
        try:
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(content_bytes)
                tmp_path = tmp.name

            try:
                import pdfplumber
                texts = []
                with pdfplumber.open(tmp_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            texts.append(page_text)
                return "\n\n".join(texts)
            except ImportError:
                pass

            try:
                import fitz
                doc = fitz.open(tmp_path)
                texts = []
                for page in doc:
                    texts.append(page.get_text())
                doc.close()
                return "\n\n".join(texts)
            except ImportError:
                pass

            return content_bytes.decode("utf-8", errors="replace")
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)

    return content_bytes.decode("utf-8", errors="replace")


_CHECK_SKILL_MAP = {
    "compliance": ("services.check.skills.compliance_check_skill", "ComplianceCheckSkill"),
    "disqualification": ("services.check.skills.disqualification_check_skill", "DisqualificationCheckSkill"),
    "qualification": ("services.check.skills.qualification_check_skill", "QualificationCheckSkill"),
    "pricing": ("services.check.skills.pricing_check_skill", "PricingCheckSkill"),
    "fitScore": ("services.check.skills.fit_score_skill", "FitScoreSkill"),
    "deposit": ("services.check.skills.deposit_check_skill", "DepositCheckSkill"),
    "signature": ("services.check.skills.signature_check_skill", "SignatureCheckSkill"),
    "validity": ("services.check.skills.validity_check_skill", "ValidityCheckSkill"),
    "consistency": ("services.check.skills.consistency_check_skill", "ConsistencyCheckSkill"),
    "duplicate": ("services.check.skills.duplicate_check_skill", "DuplicateCheckSkill"),
    "mandatoryReq": ("services.check.skills.mandatory_req_check_skill", "MandatoryReqCheckSkill"),
    "docIntegrity": ("services.check.skills.doc_integrity_check_skill", "DocIntegrityCheckSkill"),
    "aiTextCheck": ("services.check.skills.ai_text_check_skill", "AITextCheckSkill"),
    "riskScore": ("services.check.skills.risk_score_skill", "RiskScoreSkill"),
    "crossCheck": ("services.check.skills.cross_check_skill", "CrossCheckSkill"),
    "sampleReport": ("services.check.skills.sample_report_check_skill", "SampleReportCheckSkill"),
    "jointBid": ("services.check.skills.joint_bid_check_skill", "JointBidCheckSkill"),
    "ebidSubmit": ("services.check.skills.ebid_submit_check_skill", "EbidSubmitCheckSkill"),
    "pricingLogic": ("services.check.skills.pricing_logic_check_skill", "PricingLogicCheckSkill"),
}

# ---------------------------------------------------------------------------
# 全面检查：项目模式与上传模式共用的类型清单、参数装配与 skill 执行
# ---------------------------------------------------------------------------

# 全面检查覆盖的 15 项。两种模式共用一份清单，避免各自维护而漂移。
_FULL_CHECK_TYPES = [
    "compliance", "disqualification", "qualification", "pricing",
    "fitScore", "deposit", "signature", "validity",
    "consistency", "duplicate", "mandatoryReq", "docIntegrity",
    "aiTextCheck", "crossCheck", "pricingLogic",
]


def _full_check_params(
    check_type: str,
    tender_text: str,
    bid_text: str,
    *,
    bid_deadline: str = "",
    max_price: float | None = None,
    project_facts: dict | None = None,
    reference_texts: list | None = None,
) -> dict:
    """按检查类型装配 skill 入参。

    上传模式没有项目上下文，只给两份文本即可（补充参数走默认值）；
    项目模式从 Analysis 取出 bid_deadline / max_price / project_facts 再传进来。
    """
    if check_type == "duplicate":
        return {"bid_text": bid_text, "reference_texts": reference_texts or []}
    if check_type == "aiTextCheck":
        return {"bid_text": bid_text}

    params: dict = {"tender_text": tender_text, "bid_text": bid_text}
    if check_type == "qualification":
        params["bid_deadline"] = bid_deadline
    elif check_type == "pricing":
        params["max_price"] = max_price
    elif check_type == "consistency":
        params["project_facts"] = project_facts or {}
    return params


async def _exec_check_skill(
    check_type: str,
    params: dict,
    gateway,
    project_id: str = "",
) -> tuple[str, dict]:
    """加载并执行一个检查 skill，返回 (check_type, {success, data, error})。

    所有调用点统一走这里，不要再在函数体内重复写 importlib 加载逻辑：
    上传模式曾自己写过一份，而同一函数靠后的一句 `import importlib` 使
    importlib 成为该函数的局部名、遮蔽了模块级导入，嵌套闭包于是把它当作
    外层自由变量读取；fullCheck 分支永远走不到那句 import，15 项检查全部抛
    "cannot access free variable 'importlib'"，再被兜底 except 吞成 success=False，
    前端表现为「全面检查全是异常」。

    每个 skill 用自己的 DB 会话，避免并发共享同一会话。
    """
    skill_info = _CHECK_SKILL_MAP.get(check_type)
    if not skill_info:
        return check_type, {
            "success": False,
            "data": {},
            "error": f"未知的检查类型: {check_type}",
            "warnings": [],
        }

    module_path, class_name = skill_info
    try:
        module = importlib.import_module(module_path)
        skill_class = getattr(module, class_name)
        skill = skill_class()

        from services.database import async_session

        session_factory = async_session()
        async with session_factory() as skill_db:
            ctx = SkillContext(
                project_id=project_id, db=skill_db, llm=gateway, parameters=params
            )
            result = await skill.safe_execute(ctx)
        return check_type, {
            "success": result.success,
            "data": result.data,
            "error": result.error,
            "warnings": result.warnings,
        }
    except Exception as e:
        return check_type, {
            "success": False,
            "data": {},
            "error": str(e),
            "warnings": [],
        }


# ---------------------------------------------------------------------------
# Async single-check: mapping, helpers, worker, endpoint
# ---------------------------------------------------------------------------

_CHECK_TYPE_TO_ENUM: dict[str, CheckType] = {
    "compliance": CheckType.COMPLIANCE,
    "disqualification": CheckType.DISQUALIFICATION,
    "qualification": CheckType.QUALIFICATION,
    "pricing": CheckType.PRICING,
    "fitScore": CheckType.FIT_SCORE,
    "deposit": CheckType.DEPOSIT,
    "signature": CheckType.SIGNATURE,
    "validity": CheckType.VALIDITY,
    "consistency": CheckType.CONSISTENCY,
    "duplicate": CheckType.DUPLICATE,
    "mandatoryReq": CheckType.MANDATORY,
    "docIntegrity": CheckType.DOC_INTEGRITY,
    "aiTextCheck": CheckType.AI_TEXT,
    "riskScore": CheckType.RISK_SCORE,
    "crossCheck": CheckType.CROSS_CHECK,
    "sampleReport": CheckType.SAMPLE_REPORT,
    "jointBid": CheckType.JOINT_BID,
    "ebidSubmit": CheckType.EBID_SUBMIT,
    "pricingLogic": CheckType.PRICING_LOGIC,
    "selfcheck": CheckType.SELFCHECK,
}


class SingleCheckRequest(BaseModel):
    check_type: str


async def _build_check_params(
    check_type: str,
    project_id: str,
    db: AsyncSession,
    project: Project,
    tender_text: str,
    bid_text: str,
) -> dict:
    """Build skill-specific parameters based on check_type."""
    base = {"tender_text": tender_text, "bid_text": bid_text}

    if check_type in ("compliance", "fitScore", "deposit", "signature",
                      "validity", "mandatoryReq", "docIntegrity",
                      "crossCheck", "sampleReport", "jointBid",
                      "ebidSubmit", "pricingLogic", "disqualification"):
        # disqualification needs extra clauses
        if check_type == "disqualification":
            analysis_result = await db.execute(
                select(Analysis).where(Analysis.project_id == project.id)
            )
            analysis = analysis_result.scalar_one_or_none()
            dq_clauses: list = []
            if analysis and analysis.dimensions:
                dq_data = analysis.dimensions.get("disqualification", {})
                if dq_data and not dq_data.get("error"):
                    dq_clauses = dq_data.get("clauses", [])
            return {**base, "disqualification_clauses": dq_clauses}
        return base

    if check_type == "qualification":
        analysis_result = await db.execute(
            select(Analysis).where(Analysis.project_id == project.id)
        )
        analysis = analysis_result.scalar_one_or_none()
        timeline = analysis.dimensions.get("timeline", {}) if analysis and analysis.dimensions else {}
        bid_deadline = timeline.get("投标截止日", "") if isinstance(timeline, dict) else ""
        return {**base, "bid_deadline": bid_deadline}

    if check_type == "pricing":
        analysis_result = await db.execute(
            select(Analysis).where(Analysis.project_id == project.id)
        )
        analysis = analysis_result.scalar_one_or_none()
        max_price = None
        if analysis and analysis.dimensions:
            project_info = analysis.dimensions.get("project_info", {})
            if isinstance(project_info, dict):
                budget_str = project_info.get("预算金额", "")
                if budget_str:
                    match = re.search(r"[\d.]+", str(budget_str))
                    if match:
                        max_price = float(match.group())
        return {**base, "max_price": max_price}

    if check_type == "consistency":
        analysis_result = await db.execute(
            select(Analysis).where(Analysis.project_id == project.id)
        )
        analysis = analysis_result.scalar_one_or_none()
        project_facts: dict = {}
        if analysis and analysis.dimensions:
            project_facts = {"project_name": project.name, "dimensions": analysis.dimensions}
        return {**base, "project_facts": project_facts}

    if check_type == "duplicate":
        reference_texts: list[str] = []
        try:
            other_docs_result = await db.execute(
                select(Document).where(
                    Document.project_id != project.id,
                    Document.type == "bid",
                    Document.parsed_content.isnot(None),
                ).limit(10)
            )
            other_docs = other_docs_result.scalars().all()
            reference_texts = [doc.parsed_content for doc in other_docs if doc.parsed_content]
        except Exception:
            pass
        return {"bid_text": bid_text, "reference_texts": reference_texts}

    if check_type == "aiTextCheck":
        return {"bid_text": bid_text}

    if check_type == "riskScore":
        reports_result = await db.execute(
            select(CheckReport).where(CheckReport.project_id == project.id)
        )
        reports = reports_result.scalars().all()
        check_results = {
            r.type.value if isinstance(r.type, CheckType) else str(r.type): r.results
            for r in reports
        }
        return {"check_results": check_results}

    # fallback
    return base


async def _do_single_check(project_id: str, check_type: str):
    """Background worker: run a single check (or selfcheck) and persist results."""
    from services.database import async_session

    session_factory = async_session()
    async with session_factory() as db:
        project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)

        if not tender_text or not bid_text:
            raise ValueError("招标文件或投标文件内容为空")

        gateway = await get_agent_gateway(db, "check")

        # --- selfcheck: 3 sub-skills in parallel + summary skill ---
        if check_type == "selfcheck":
            analysis_result = await db.execute(
                select(Analysis).where(Analysis.project_id == project.id)
            )
            analysis = analysis_result.scalar_one_or_none()
            max_price = None
            bid_deadline = ""
            if analysis and analysis.dimensions:
                pi = analysis.dimensions.get("project_info", {})
                if isinstance(pi, dict):
                    budget_str = pi.get("预算金额", "")
                    if budget_str:
                        m = re.search(r"[\d.]+", str(budget_str))
                        if m:
                            max_price = float(m.group())
                timeline = analysis.dimensions.get("timeline", {})
                if isinstance(timeline, dict):
                    bid_deadline = timeline.get("投标截止日", "")

            check_results: dict = {}

            async def _run_sub(skill_cls, params: dict, key: str):
                skill = skill_cls()
                async with session_factory() as sub_db:
                    ctx = SkillContext(project_id=project_id, db=sub_db, llm=gateway, parameters=params)
                    res = await skill.safe_execute(ctx)
                return key, res

            from services.check.skills.compliance_check_skill import ComplianceCheckSkill
            from services.check.skills.pricing_check_skill import PricingCheckSkill
            from services.check.skills.qualification_check_skill import QualificationCheckSkill

            sub_tasks = [
                _run_sub(ComplianceCheckSkill, {"tender_text": tender_text, "bid_text": bid_text}, "compliance_check"),
                _run_sub(PricingCheckSkill, {"tender_text": tender_text, "bid_text": bid_text, "max_price": max_price}, "pricing_check"),
                _run_sub(QualificationCheckSkill, {"tender_text": tender_text, "bid_text": bid_text, "bid_deadline": bid_deadline}, "qualification_check"),
            ]
            sub_results = await asyncio.gather(*sub_tasks, return_exceptions=True)
            for sr in sub_results:
                if isinstance(sr, Exception):
                    logger.warning(f"[selfcheck] sub-skill error: {sr}")
                    continue
                key, res = sr
                if res.success:
                    check_results[key] = res.data

            from services.check.skills.selfcheck_list_skill import SelfcheckListSkill
            skill = SelfcheckListSkill()
            async with session_factory() as skill_db:
                ctx = SkillContext(
                    project_id=project_id, db=skill_db, llm=gateway,
                    parameters={"check_results": check_results},
                )
                skill_result = await skill.safe_execute(ctx)

            if skill_result.success:
                report = CheckReport(
                    project_id=project.id,
                    type=CheckType.SELFCHECK,
                    results=skill_result.data,
                    risk_level="low" if skill_result.data.get("can_submit") else "high",
                )
                db.add(report)
                await db.flush()
                await db.commit()

            return {
                "success": skill_result.success,
                "data": skill_result.data,
                "error": skill_result.error,
                "warnings": skill_result.warnings,
            }

        # --- normal single check ---
        params = await _build_check_params(check_type, project_id, db, project, tender_text, bid_text)
        skill_info = _CHECK_SKILL_MAP.get(check_type)
        if not skill_info:
            raise ValueError(f"未知的检查类型: {check_type}")

        module_path, class_name = skill_info
        module = importlib.import_module(module_path)
        skill_class = getattr(module, class_name)
        skill = skill_class()

        async with session_factory() as skill_db:
            ctx = SkillContext(project_id=project_id, db=skill_db, llm=gateway, parameters=params)
            skill_result = await skill.safe_execute(ctx)

        if skill_result.success:
            check_enum = _CHECK_TYPE_TO_ENUM.get(check_type)
            if check_enum:
                risk_level = "low"
                if isinstance(skill_result.data, dict):
                    risk_level = (
                        skill_result.data.get("risk_level")
                        or skill_result.data.get("overall_risk")
                        or "low"
                    )
                report = CheckReport(
                    project_id=project.id,
                    type=check_enum,
                    results=skill_result.data,
                    risk_level=risk_level,
                )
                db.add(report)
                await db.flush()
                await db.commit()

        return {
            "success": skill_result.success,
            "data": skill_result.data,
            "error": skill_result.error,
            "warnings": skill_result.warnings,
        }


@router.post("/{project_id}/check-async")
async def submit_single_check(
    project_id: str,
    body: SingleCheckRequest,
    db: AsyncSession = Depends(get_db),
):
    """Submit any single check as an async task. Returns task_id for polling."""
    check_type = body.check_type

    # Validate project
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    # Validate check_type
    valid_types = set(_CHECK_TYPE_TO_ENUM.keys())
    if check_type not in valid_types:
        raise HTTPException(
            status_code=400,
            detail=f"不支持的检查类型: {check_type}。支持: {', '.join(sorted(valid_types))}",
        )

    tm = TaskManager.instance()
    task = await tm.submit("single_check", _do_single_check, project_id, check_type)

    return {
        "task_id": task.task_id,
        "status": "pending",
        "message": f"{check_type} 检查任务已提交",
    }



@router.post("/tender-bid-review")
async def tender_bid_review(
    bid_file: UploadFile = File(..., description="投标书(.docx/.pdf/.txt)"),
    tender_file: UploadFile = File(..., description="招标文件(.docx/.pdf/.txt)"),
    company_name: str = Form("", description="投标方公司名称"),
    school_name: str = Form("", description="招标方学校名称"),
    db: AsyncSession = Depends(get_db),
):
    """投标文件审查：招标文件 + 投标书六维度交叉审查，返回异步任务ID。

    任务完成后 task.result 包含 excel_base64 / file_name，前端可直接下载。
    """
    bid_text = await _parse_review_document(bid_file)
    if not bid_text.strip():
        raise HTTPException(status_code=400, detail="投标书内容为空或无法解析")

    tender_text = await _parse_review_document(tender_file)
    if not tender_text.strip():
        raise HTTPException(status_code=400, detail="招标文件内容为空或无法解析")

    bid_filename = bid_file.filename or "bid"
    tender_filename = tender_file.filename if tender_file else None

    tm = TaskManager.instance()
    task = await tm.submit(
        "tender_bid_review",
        _do_tender_bid_review,
        tender_text,
        bid_text,
        company_name,
        school_name,
        bid_filename,
        tender_filename,
    )

    return {
        "task_id": task.task_id,
        "status": "pending",
        "source": "upload",
        "bid_filename": bid_filename,
        "tender_filename": tender_filename,
        "message": "投标文件审查任务已提交，请通过 GET /check/task/{task_id} 查询进度",
    }


async def _do_tender_bid_review(
    tender_text: str,
    bid_text: str,
    company_name: str,
    school_name: str,
    bid_filename: str,
    tender_filename: str | None,
):
    """后台执行投标文件审查并生成 Excel。必须自建 DB 会话。"""
    from services.database import async_session
    from services.check.skills.tender_bid_review_skill import TenderBidReviewSkill

    session_factory = async_session()
    async with session_factory() as skill_db:
        gateway = await get_agent_gateway(skill_db, "check")
        skill = TenderBidReviewSkill()
        ctx = SkillContext(
            project_id="",
            db=skill_db,
            llm=gateway,
            progress_callback=functools.partial(
                _report_review_progress,
                task.task_id,
            ),
            parameters={
                "tender_lines": tender_text,
                "bid_lines": bid_text,
                "company_name": company_name,
                "school_name": school_name,
            },
        )
        result = await skill.safe_execute(ctx)

        response = {
            "success": result.success,
            "data": result.data,
            "error": result.error,
            "warnings": result.warnings,
            "source": "upload",
            "bid_filename": bid_filename,
            "tender_filename": tender_filename,
        }

        if result.success and result.data.get("excel_base64"):
            import base64
            excel_bytes = base64.b64decode(result.data["excel_base64"])
            file_name = result.data.get("file_name", "投标文件审查.xlsx")
            import os
            import tempfile
            temp_dir = os.path.join(tempfile.gettempdir(), "bidmaster_exports")
            os.makedirs(temp_dir, exist_ok=True)
            safe_name = re.sub(r'[\\/:*?\"<>|]', '_', file_name).strip()
            safe_name = re.sub(r'_+', '_', safe_name)
            if not safe_name.endswith('.xlsx'):
                safe_name += '.xlsx'
            out_path = os.path.join(temp_dir, safe_name)
            with open(out_path, "wb") as f:
                f.write(excel_bytes)
            response["data"]["download_url"] = f"/check/tender-bid-review/download/{safe_name}"
            response["data"]["file_name"] = safe_name

        return response


async def _report_review_progress(task_id: str, skill_name: str, stage: str, payload: dict) -> None:
    stage_messages = {
        "started": "任务已启动",
        "dimension_started": "正在生成审查任务",
        "chunk_started": "正在交叉审查",
        "chunk_completed": "已完成审查片段",
        "dimension_completed": "已完成审查维度",
    }
    if stage == "completed":
        TaskManager.instance().set_progress(task_id, 1.0, "审查完成")
        return
    if stage == "failed":
        TaskManager.instance().set_progress(task_id, float(payload.get("progress", 0.0)), "审查失败")
        return
    TaskManager.instance().set_progress(
        task_id,
        float(payload.get("progress", 0.0)),
        f"{stage_messages.get(stage, stage)}（{payload.get('completed', 0)}/{payload.get('total', 0)}）",
    )


async def _parse_review_document(file: UploadFile) -> str:
    """使用统一文档引擎解析审查文件，保留行号锚点并拒绝空扫描件。"""
    from core.doc_engine import get_parser

    suffix = os.path.splitext(file.filename or "")[1].lower()
    content_bytes = await file.read()
    if len(content_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    if not suffix:
        raise HTTPException(status_code=400, detail="文件缺少扩展名")

    temp_path = ""
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temp_file:
            temp_file.write(content_bytes)
            temp_path = temp_file.name
        parsed = get_parser(suffix).parse(temp_path)
        parsed.text = parsed.text.lstrip("\ufeff")
        lines = [line.strip() for line in parsed.text.splitlines() if line.strip()]
        for table in parsed.tables or []:
            rows = table.get("rows", []) if isinstance(table, dict) else []
            for row in rows:
                cells = [str(cell).strip().replace("\n", " ") for cell in row]
                if any(cells):
                    lines.append(" | ".join(cells))
        numbered = "\n".join(f"{line_no}\t{text}" for line_no, text in enumerate(lines, 1))
        if len("".join(lines)) < 50:
            raise HTTPException(status_code=400, detail="文件内容为空、过少或可能是扫描件，请上传可复制文本版")
        return numbered
    except HTTPException:
        raise
    except Exception as exc:
        logger.warning("投标审查文件解析失败 %s: %s", file.filename, exc)
        raise HTTPException(status_code=400, detail=f"文件解析失败：{file.filename}") from exc
    finally:
        if temp_path and os.path.exists(temp_path):
            os.unlink(temp_path)


@router.get("/tender-bid-review/download/{file_name}")
async def download_tender_bid_review(file_name: str):
    """下载投标文件审查 Excel。只允许下载 bidmaster_exports 目录下文件。"""
    import tempfile
    import os
    from urllib.parse import quote

    temp_dir = os.path.join(tempfile.gettempdir(), "bidmaster_exports")
    safe_name = os.path.basename(file_name)
    file_path = os.path.join(temp_dir, safe_name)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="文件不存在或已过期")
    if not safe_name.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")

    from fastapi.responses import FileResponse
    encoded = quote(safe_name)
    return FileResponse(
        path=file_path,
        filename=safe_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )

@router.post("/upload-check")
async def upload_and_check(
    bid_file: UploadFile = File(..., description="投标文件(.docx/.pdf/.txt)"),
    tender_file: UploadFile | None = File(None, description="招标文件(可选，.docx/.pdf/.txt)"),
    check_type: str = Form("fullCheck", description="检查类型: fullCheck/compliance/disqualification/..."),
    db: AsyncSession = Depends(get_db),
):
    """上传模式检查：解析完文件立即返回 task_id，前端轮询 GET /check/task/{task_id}。

    实测单项检查就要 70-80s，全面检查 15 项即便限流并发也要数分钟，同步等待必然撞上
    axios(300s) / nginx(600s) 超时——所以和项目模式 full-check 一样改成异步任务。
    文件解析仍在这里同步做，格式问题能立刻以 400 反馈，不必等轮询。

    db 只用于 get_db 的数据库就绪门禁（DB 不可用时直接 503，而不是提交一个注定失败的任务）。
    """
    if check_type not in _CHECK_SKILL_MAP and check_type not in ("fullCheck", "selfcheck"):
        raise HTTPException(status_code=400, detail=f"不支持的检查类型: {check_type}")

    bid_text = await _parse_uploaded_file(bid_file)
    if not bid_text.strip():
        raise HTTPException(status_code=400, detail="投标文件内容为空或无法解析")

    tender_text = ""
    if tender_file:
        tender_text = await _parse_uploaded_file(tender_file)

    bid_filename = bid_file.filename or "bid"
    tender_filename = tender_file.filename if tender_file else None

    tm = TaskManager.instance()
    task = await tm.submit(
        "upload_check",
        _do_upload_check,
        check_type,
        tender_text,
        bid_text,
        bid_filename,
        tender_filename,
    )

    return {
        "task_id": task.task_id,
        "status": "pending",
        "source": "upload",
        "check_type": check_type,
        "bid_filename": bid_filename,
        "tender_filename": tender_filename,
        "message": (
            f"全面检查任务已提交（15 项，通常需 3-10 分钟），请通过 GET /check/task/{task.task_id} 查询进度"
            if check_type == "fullCheck"
            else f"{check_type} 检查任务已提交，请通过 GET /check/task/{task.task_id} 查询进度"
        ),
    }


async def _do_upload_check(
    check_type: str,
    tender_text: str,
    bid_text: str,
    bid_filename: str,
    tender_filename: str | None,
):
    """后台执行上传模式检查。必须自建 DB 会话：请求级会话在响应返回后就关闭了。"""
    from services.database import async_session

    source_info = {
        "source": "upload",
        "bid_filename": bid_filename,
        "tender_filename": tender_filename,
    }

    session_factory = async_session()
    async with session_factory() as db:
        gateway = await get_agent_gateway(db, "check")

        if check_type == "fullCheck":
            semaphore = asyncio.Semaphore(_check_max_concurrent())

            async def _run(ct: str) -> tuple[str, dict]:
                params = _full_check_params(ct, tender_text, bid_text)
                async with semaphore:
                    return await _exec_check_skill(ct, params, gateway)

            gather_results = await asyncio.gather(
                *[_run(ct) for ct in _FULL_CHECK_TYPES],
                return_exceptions=True,
            )

            all_results: dict = {}
            for r in gather_results:
                if isinstance(r, Exception):
                    logger.warning(f"[upload-check] skill exception: {r}")
                    continue
                ct, result_dict = r
                all_results[ct] = result_dict

            has_critical = any(
                r.get("data", {}).get("has_critical_issues")
                or r.get("data", {}).get("risk_level") == "high"
                for r in all_results.values()
                if r.get("success") and isinstance(r.get("data"), dict)
            )

            return {
                "success": True,
                "data": all_results,
                "has_critical": has_critical,
                **source_info,
            }

        if check_type == "selfcheck":
            # 上传模式没有前序检查结果可汇总，沿用原行为：把空清单交给 skill 自行判断
            from services.check.skills.selfcheck_list_skill import SelfcheckListSkill

            skill = SelfcheckListSkill()
            ctx = SkillContext(
                project_id="", db=db, llm=gateway, parameters={"check_results": {}}
            )
            skill_result = await skill.safe_execute(ctx)
            return {
                "success": skill_result.success,
                "data": skill_result.data,
                "error": skill_result.error,
                "warnings": skill_result.warnings,
                **source_info,
            }

        params = _full_check_params(check_type, tender_text, bid_text)
        _, result = await _exec_check_skill(check_type, params, gateway)
        return {**result, **source_info}


@router.post("/{project_id}/signature")
async def check_signature(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.signature_check_skill import SignatureCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = SignatureCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.SIGNATURE, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/validity")
async def check_validity(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.validity_check_skill import ValidityCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = ValidityCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.VALIDITY, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/consistency")
async def check_consistency(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.consistency_check_skill import ConsistencyCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = ConsistencyCheckSkill()

    # Build project_facts from analysis data for rule-based consistency checks
    analysis_result = await db.execute(
        select(Analysis).where(Analysis.project_id == project.id)
    )
    analysis = analysis_result.scalar_one_or_none()
    project_facts = {}
    if analysis and analysis.dimensions:
        project_facts = {
            "project_name": project.name,
            "dimensions": analysis.dimensions,
        }

    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text, "project_facts": project_facts})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.CONSISTENCY, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/duplicate")
async def check_duplicate(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.duplicate_check_skill import DuplicateCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = DuplicateCheckSkill()

    # Build reference_texts from other projects' bid documents for cross-project dedup
    reference_texts = []
    try:
        other_docs_result = await db.execute(
            select(Document).where(
                Document.project_id != project.id,
                Document.type == "bid",
                Document.parsed_content.isnot(None),
            ).limit(10)
        )
        other_docs = other_docs_result.scalars().all()
        reference_texts = [doc.parsed_content for doc in other_docs if doc.parsed_content]
    except Exception:
        pass

    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"bid_text": bid_text, "reference_texts": reference_texts})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.DUPLICATE, results=skill_result.data, risk_level=skill_result.data.get("overall_risk", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/mandatory-req")
async def check_mandatory_req(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    if not tender_text or not bid_text:
        raise HTTPException(status_code=400, detail="招标文件或投标文件内容为空")
    from services.check.skills.mandatory_req_check_skill import MandatoryReqCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = MandatoryReqCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.MANDATORY, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.get("/{project_id}/reports/{report_id}/export")
async def export_check_report(
    project_id: str,
    report_id: str,
    format: str = "markdown",
    db: AsyncSession = Depends(get_db),
):
    from fastapi.responses import PlainTextResponse

    result = await db.execute(select(CheckReport).where(CheckReport.id == report_id))
    report = result.scalar_one_or_none()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")

    project_result = await db.execute(select(Project).where(Project.id == project_id))
    project = project_result.scalar_one_or_none()
    project_name = project.name if project else "未命名项目"

    from services.check.skills.check_report_export_skill import CheckReportExportSkill
    from services.llm_factory import get_agent_gateway

    gateway = await get_agent_gateway(db, "check")
    skill = CheckReportExportSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={
            "report_data": report.results or {},
            "format": format,
            "project_name": project_name,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if not skill_result.success:
        raise HTTPException(status_code=500, detail=skill_result.error)

    content = skill_result.data.get("content", "")
    content_type = "text/markdown" if format == "markdown" else "text/html" if format == "html" else "application/json"
    return PlainTextResponse(content=content, media_type=content_type)


@router.get("/{project_id}/reports/{report_id}/content")
async def get_check_report_content(
    project_id: str,
    report_id: str,
    format: str = "markdown",
    db: AsyncSession = Depends(get_db),
):
    """直接返回报告内容(用于前端预览)，不触发下载。"""
    from services.check.skills.check_report_export_skill import CheckReportExportSkill
    from services.llm_factory import get_agent_gateway

    result = await db.execute(select(CheckReport).where(CheckReport.id == report_id))
    report = result.scalar_one_or_none()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")

    project_result = await db.execute(select(Project).where(Project.id == project_id))
    project = project_result.scalar_one_or_none()
    project_name = project.name if project else "未命名项目"

    if not report.results:
        return {
            "success": True,
            "report_id": report_id,
            "format": format,
            "content": "",
            "project_name": project_name,
            "size": 0,
        }

    gateway = await get_agent_gateway(db, "check")
    skill = CheckReportExportSkill()
    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={
            "report_data": report.results,
            "format": format,
            "project_name": project_name,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if not skill_result.success:
        raise HTTPException(status_code=500, detail=skill_result.error or "报告生成失败")

    return {
        "success": True,
        "report_id": report_id,
        "format": format,
        "content": skill_result.data.get("content", ""),
        "size": skill_result.data.get("size", 0),
        "project_name": project_name,
        "generated_at": skill_result.data.get("generated_at"),
    }


@router.post("/{project_id}/doc-integrity")
async def check_doc_integrity(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.doc_integrity_check_skill import DocIntegrityCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = DocIntegrityCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.DOC_INTEGRITY, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/ai-text-check")
async def check_ai_text(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.ai_text_check_skill import AITextCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = AITextCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.AI_TEXT, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/risk-score")
async def check_risk_score(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    reports_result = await db.execute(
        select(CheckReport).where(CheckReport.project_id == project.id)
    )
    reports = reports_result.scalars().all()
    check_results = {r.type.value if isinstance(r.type, CheckType) else str(r.type): r.results for r in reports}
    from services.check.skills.risk_score_skill import RiskScoreSkill
    skill = RiskScoreSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=await get_agent_gateway(db, "check"), parameters={"check_results": check_results})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(
            project_id=project.id,
            type=CheckType.RISK_SCORE,
            results=skill_result.data,
            risk_level=skill_result.data.get("risk_level", "low"),
            summary={"score": skill_result.data.get("score", 0)},
        )
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error}


@router.post("/{project_id}/cross-check")
async def check_cross(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.cross_check_skill import CrossCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = CrossCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.CROSS_CHECK, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/sample-report")
async def check_sample_report(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.sample_report_check_skill import SampleReportCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = SampleReportCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.SAMPLE_REPORT, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/joint-bid")
async def check_joint_bid(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.joint_bid_check_skill import JointBidCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = JointBidCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.JOINT_BID, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/ebid-submit")
async def check_ebid_submit(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.ebid_submit_check_skill import EbidSubmitCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = EbidSubmitCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.EBID_SUBMIT, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}


@router.post("/{project_id}/pricing-logic")
async def check_pricing_logic(project_id: str, db: AsyncSession = Depends(get_db)):
    project, tender_text, bid_text = await _get_tender_and_bid_text(project_id, db)
    from services.check.skills.pricing_logic_check_skill import PricingLogicCheckSkill
    gateway = await get_agent_gateway(db, "check")
    skill = PricingLogicCheckSkill()
    ctx = SkillContext(project_id=project_id, db=db, llm=gateway, parameters={"tender_text": tender_text, "bid_text": bid_text})
    skill_result = await skill.safe_execute(ctx)
    if skill_result.success:
        report = CheckReport(project_id=project.id, type=CheckType.PRICING_LOGIC, results=skill_result.data, risk_level=skill_result.data.get("risk_level", "low"))
        db.add(report)
        await db.flush()
    return {"success": skill_result.success, "data": skill_result.data, "error": skill_result.error, "warnings": skill_result.warnings}
