from __future__ import annotations

import io
import logging
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from core.http_headers import content_disposition
from core.settings import get_settings
from services.database import get_db
from services.llm_factory import get_agent_gateway
from core.skill_engine.base import SkillContext

logger = logging.getLogger(__name__)

router = APIRouter()

UPLOAD_DIR = Path("./uploads/formatted")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

TEMPLATE_DIR = Path("./templates")
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)

EXPORT_DIR = Path("./uploads/exports")
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

MAX_UPLOAD_BYTES = 50 * 1024 * 1024


def _save_upload(file: UploadFile, prefix: str = "") -> Path:
    if file.size and file.size > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    name = Path(file.filename or "upload").name
    safe = "".join(c for c in name if c.isalnum() or c in ("-", "_", "."))
    dest = UPLOAD_DIR / f"{prefix}{safe}"
    content = file.file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    dest.write_bytes(content)
    return dest


def _unlink_quietly(*paths: Path) -> None:
    """删掉导出过程中的临时落盘文件。数据此时已 read_bytes() 进内存，删失败也不影响响应。"""
    for path in paths:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass


@router.post("/format")
async def format_document(
    file: UploadFile = File(...),
    template: str = "default",
    mode: str = Query("format", description="format|check|diff|beautify"),
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file)

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = await get_agent_gateway(db, "format")
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "template": template,
            "mode": mode,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success and skill_result.data:
        result_data = skill_result.data
        output_path = result_data.get("output_path", "")
        if output_path and Path(output_path).exists():
            result_data["file_name"] = Path(output_path).name
            try:
                result_data["file_size"] = Path(output_path).stat().st_size
            except OSError:
                pass

        return {"success": True, "data": result_data}

    return {"success": False, "error": skill_result.error}


@router.post("/check-format")
async def check_format(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="check_")

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = await get_agent_gateway(db, "format")
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "template": template,
            "mode": "check",
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        return {"success": True, "data": skill_result.data}
    return {"success": False, "error": skill_result.error}


@router.post("/diff-format")
async def diff_format(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="diff_")

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = await get_agent_gateway(db, "format")
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "template": template,
            "mode": "diff",
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        return {"success": True, "data": skill_result.data}
    return {"success": False, "error": skill_result.error}


@router.post("/beautify")
async def beautify_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="beautify_")

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = await get_agent_gateway(db, "format")
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "mode": "beautify",
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success and skill_result.data:
        return {"success": True, "data": skill_result.data}
    return {"success": False, "error": skill_result.error}


@router.post("/export-doc")
async def export_to_doc(
    file: UploadFile = File(...),
    # 必须显式 Form：前端把 template 放在 multipart body 里（api.ts exportDoc），
    # 不标注时 FastAPI 按 query 参数解析，body 里的值被静默丢弃 ——
    # 用户在「一键排版」选的模板对三个导出永远不生效，恒用 default。
    template: str = Form("default"),
    apply_format: bool = Form(True),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版(可选) → 转 doc 字节流下载。"""
    src = _save_upload(file, prefix="srcdocx_")

    docx_to_send: Path = src
    if apply_format:
        formatted = await _run_format_skill(src, template, db)
        if formatted:
            docx_to_send = formatted

    from services.format.skills.doc_export_skill import DocExportSkill
    skill = DocExportSkill()
    if not docx_to_send.exists():
        _unlink_quietly(src)
        raise HTTPException(status_code=500, detail="排版产物丢失")

    result = skill._try_libreoffice_doc(docx_to_send, docx_to_send.parent)
    if not result.get("success"):
        _unlink_quietly(docx_to_send, src)
        raise HTTPException(status_code=500, detail=result.get("error", "doc 转换失败"))

    doc_path = Path(result["output_path"])
    data = doc_path.read_bytes()
    # 此前只删排版产物、从不删 _save_upload 落下的 srcdocx_*，
    # 每点一次导出就在 uploads/formatted 留一份源文件副本（大标书 1.2MB/次）
    _unlink_quietly(doc_path, docx_to_send, src)

    filename = f"{Path(file.filename or 'document').stem}.doc"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/msword",
        headers=content_disposition(filename),
    )


@router.post("/export-pdf")
async def export_to_pdf(
    file: UploadFile = File(...),
    template: str = Form("default"),
    apply_format: bool = Form(True),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版(可选) → 转 pdf 字节流下载。"""
    src = _save_upload(file, prefix="srcdocx_")

    docx_to_send: Path = src
    if apply_format:
        formatted = await _run_format_skill(src, template, db)
        if formatted:
            docx_to_send = formatted

    from services.format.skills.pdf_export_skill import PdfExportSkill
    gateway = await get_agent_gateway(db, "format")
    skill = PdfExportSkill()
    out_dir = docx_to_send.parent
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "input_path": str(docx_to_send),
            "output_dir": str(out_dir),
        },
    )
    skill_result = await skill.safe_execute(ctx)
    if not skill_result.success or not skill_result.data:
        _unlink_quietly(docx_to_send, src)
        raise HTTPException(status_code=500, detail=skill_result.error or "PDF 转换失败")

    pdf_path = Path(skill_result.data["output_path"])
    data = pdf_path.read_bytes()
    _unlink_quietly(pdf_path, docx_to_send, src)

    filename = f"{Path(file.filename or 'document').stem}.pdf"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/pdf",
        headers=content_disposition(filename),
    )


@router.post("/export-formatted-docx")
async def export_formatted_docx(
    file: UploadFile = File(...),
    template: str = Form("default"),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版 → 直接下载 docx 字节流。"""
    src = _save_upload(file, prefix="srcdocx_")
    formatted = await _run_format_skill(src, template, db) or src
    data = formatted.read_bytes()
    _unlink_quietly(formatted, src)
    filename = f"{Path(file.filename or 'document').stem}_formatted.docx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=content_disposition(filename),
    )


async def _run_format_skill(src: Path, template: str, db: AsyncSession) -> Path | None:
    from services.format.skills.docx_format_skill import DocxFormatSkill
    gateway = await get_agent_gateway(db, "format")
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(src),
            "template": template,
            "mode": "format",
        },
    )
    result = await skill.safe_execute(ctx)
    if not result.success or not result.data:
        logger.warning("format skill 失败: %s", result.error)
        return None
    out = result.data.get("output_path", "")
    if out and Path(out).exists():
        return Path(out)
    return None


@router.post("/from-project/{project_id}")
async def format_from_project(
    project_id: str,
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    """从项目章节组装为 docx 文件（无需上传）"""
    from sqlalchemy import select
    from services.models import Project, Chapter

    proj_result = await db.execute(select(Project).where(Project.id == project_id))
    project = proj_result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    ch_result = await db.execute(
        select(Chapter).where(Chapter.project_id == project_id).order_by(Chapter.sort_order)
    )
    chapters = ch_result.scalars().all()
    if not chapters:
        raise HTTPException(
            status_code=400,
            detail="项目尚无任何章节大纲，请先到「投标生成」完成大纲生成。",
        )

    empty_chapters = [ch for ch in chapters if not (ch.content and ch.content.strip())]
    if empty_chapters:
        empty_titles = [ch.title for ch in empty_chapters[:5]]
        detail_msg = (
            f"项目存在 {len(empty_chapters)}/{len(chapters)} 个章节尚未生成正文，"
            f"无法执行文档输出。请先到「投标生成」完成正文生成。"
            f"未生成章节示例: {', '.join(empty_titles)}"
        )
        if len(empty_chapters) == len(chapters):
            raise HTTPException(status_code=400, detail=detail_msg)
        if len(empty_chapters) / len(chapters) > 0.5:
            raise HTTPException(
                status_code=400,
                detail=detail_msg + "（空内容章节超过50%，已禁止操作）",
            )

    title = f"{project.name} - 投标文件"

    from docx import Document
    from docx.shared import Pt
    from pathlib import Path
    from core.skill_engine.base import SkillContext
    from services.format.skills.docx_format_skill import DocxFormatSkill

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    h1 = doc.add_heading(title, level=0)

    for ch in chapters:
        if not (ch.content and ch.content.strip()):
            continue
        doc.add_heading(ch.title, level=1)
        for line in ch.content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            else:
                doc.add_paragraph(line)

    output_dir = UPLOAD_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[^\w\u4e00-\u9fff\-_]', '_', project.name)[:60] or "bid"
    src_path = output_dir / f"{safe_name}_src_{project_id[:8]}.docx"
    doc.save(str(src_path))

    from services.llm_factory import get_agent_gateway
    gateway = await get_agent_gateway(db, "format")
    skill = DocxFormatSkill()

    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={"file_path": str(src_path), "template": template, "mode": "format"},
    )
    result = await skill.safe_execute(ctx)
    if not result.success:
        try:
            src_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=result.error or "项目章节组装失败")

    final_output = result.data.get("output_path") if isinstance(result.data, dict) else None
    if final_output and Path(final_output) != src_path:
        try:
            src_path.unlink(missing_ok=True)
        except OSError:
            pass

    return {
        "success": True,
        "output_path": final_output,
        "project_name": project.name,
        "chapter_count": len(chapters),
        "generated_chapter_count": len(chapters) - len(empty_chapters),
        "empty_chapter_count": len(empty_chapters),
        "chapters": [
            {
                "id": str(c.id),
                "title": c.title,
                "word_count": c.word_count or 0,
                "status": c.status,
                "has_content": bool(c.content and c.content.strip()),
            }
            for c in chapters
        ],
    }


@router.get("/templates")
async def list_templates():
    templates = [
        {"name": "default", "description": "默认标书排版模板(仿宋正文+黑体标题)"},
        {"name": "government", "description": "政府采购标书模板(方正小标宋+仿宋)"},
        {"name": "engineering", "description": "工程标书模板(宋体正文+黑体标题)"},
    ]

    for tpl_file in TEMPLATE_DIR.glob("*.yaml"):
        name = tpl_file.stem
        if not any(t["name"] == name for t in templates):
            templates.append({"name": name, "description": f"自定义模板: {name}"})

    return {"templates": templates}


@router.get("/templates/{template_name}")
async def get_template_detail(template_name: str):
    from services.format.skills.docx_format_skill import DEFAULT_FORMAT_CONFIG

    template_path = TEMPLATE_DIR / f"{template_name}.yaml"
    if template_path.exists():
        try:
            import yaml
            with open(template_path, encoding="utf-8") as f:
                custom_config = yaml.safe_load(f) or {}
            merged = {**DEFAULT_FORMAT_CONFIG, **custom_config}
        except Exception:
            merged = DEFAULT_FORMAT_CONFIG
    else:
        merged = DEFAULT_FORMAT_CONFIG

    return {"name": template_name, "config": merged}


@router.put("/templates/{template_name}")
async def save_template(template_name: str, config: dict):
    import yaml

    template_path = TEMPLATE_DIR / f"{template_name}.yaml"
    with open(template_path, "w", encoding="utf-8") as f:
        yaml.dump(config, f, allow_unicode=True, default_flow_style=False)

    return {"success": True, "name": template_name}


@router.delete("/templates/{template_name}")
async def delete_template(template_name: str):
    if template_name in ("default", "government", "engineering"):
        return {"success": False, "error": "内置模板不可删除"}

    template_path = TEMPLATE_DIR / f"{template_name}.yaml"
    if template_path.exists():
        template_path.unlink()
        return {"success": True}
    return {"success": False, "error": "模板不存在"}


@router.get("/download")
async def download_formatted_file(path: str):
    """下载由 /format 产生的产物（限定 uploads/formatted 目录）。"""
    root = UPLOAD_DIR.resolve()
    raw = Path(path)
    # 调用方既可能传裸文件名，也可能传已含 uploads/formatted 前缀的相对路径
    # （skill 返回的 output_path 就是后者），无条件拼 UPLOAD_DIR 会把前缀重复一次。
    candidates = [raw.resolve()] if raw.is_absolute() else [(UPLOAD_DIR / path).resolve(), raw.resolve()]
    inside = [c for c in candidates if c.is_relative_to(root)]
    if not inside:
        raise HTTPException(status_code=400, detail="非法路径")
    p = next((c for c in inside if c.is_file()), None)
    if p is None:
        raise HTTPException(status_code=404, detail="文件不存在")
    return StreamingResponse(
        p.open("rb"),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=content_disposition(p.name),
    )
