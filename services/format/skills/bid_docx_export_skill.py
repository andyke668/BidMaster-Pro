from __future__ import annotations

import io
import re
import logging
from pathlib import Path

from core.skill_engine.base import Skill, SkillContext, SkillResult

logger = logging.getLogger(__name__)


class BidDocxExportSkill(Skill):
    name = "bid_docx_export"
    description = "投标文件Word导出(大纲结构+正文+封面+占位符填充+排版)"
    category = "format"
    version = "1.0.0"
    triggers = ["导出Word", "生成Word", "导出投标文件"]

    async def execute(self, ctx: SkillContext) -> SkillResult:
        outline_tree = ctx.parameters.get("outline_tree", {})
        chapters_content = ctx.parameters.get("chapters_content", {})
        project_facts = ctx.parameters.get("project_facts", {})
        cover_data = ctx.parameters.get("cover_data", {})
        format_config = ctx.parameters.get("format_config", {})

        if not outline_tree:
            return SkillResult(success=False, error="大纲数据为空")

        try:
            doc = self._build_document(outline_tree, chapters_content, project_facts, cover_data, format_config)
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            docx_bytes = buffer.read()

            return SkillResult(
                success=True,
                data={
                    "docx_bytes": docx_bytes,
                    "file_size": len(docx_bytes),
                    "format": "docx",
                },
            )
        except Exception as e:
            logger.error(f"Word导出失败: {e}")
            return SkillResult(success=False, error=str(e))

    def _build_document(self, outline_tree, chapters_content, project_facts, cover_data, format_config):
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document()
        config = self._merge_config(format_config)

        self._set_page_format(doc, config)
        self._set_styles(doc, config)

        project_name = project_facts.get("project_name", "")
        bidder_name = project_facts.get("bidder_name", project_facts.get("company_name", ""))

        self._generate_cover(doc, project_name, bidder_name, cover_data, config)

        chapters = outline_tree.get("chapters", []) if isinstance(outline_tree, dict) else outline_tree
        self._add_chapters(doc, chapters, chapters_content, project_facts, config)

        self._add_page_numbers(doc)

        return doc

    def _merge_config(self, user_config: dict) -> dict:
        defaults = {
            "margin_top": 2.54, "margin_bottom": 2.54,
            "margin_left": 3.18, "margin_right": 2.54,
            "title_font": "方正小标宋简体",
            "h1_font": "黑体", "h2_font": "黑体", "h3_font": "黑体",
            "h4_font": "楷体_GB2312",
            "body_font": "仿宋_GB2312",
            "english_font": "Times New Roman",
            "title_size": 22, "h1_size": 16, "h2_size": 15,
            "h3_size": 15, "h4_size": 14, "body_size": 12,
            "heading_number_format": "decimal",
            "line_spacing": 28,
        }
        defaults.update(user_config)
        return defaults

    def _set_page_format(self, doc, config):
        for section in doc.sections:
            section.top_margin = Cm(config["margin_top"])
            section.bottom_margin = Cm(config["margin_bottom"])
            section.left_margin = Cm(config["margin_left"])
            section.right_margin = Cm(config["margin_right"])
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)

    def _set_styles(self, doc, config):
        from docx.shared import Pt

        style = doc.styles["Normal"]
        style.font.name = config["english_font"]
        style.font.size = Pt(config["body_size"])
        style.element.rPr.rFonts.set(qn("w:eastAsia"), config["body_font"])

        for i in range(1, 5):
            style_name = f"Heading {i}"
            if style_name in doc.styles:
                s = doc.styles[style_name]
                font_key = f"h{i}_font"
                size_key = f"h{i}_size"
                s.font.name = config.get(font_key, "黑体")
                s.font.size = Pt(config.get(size_key, 16 - i))
                s.font.bold = True
                s.font.color.rgb = RGBColor(0, 0, 0)
                s.element.rPr.rFonts.set(qn("w:eastAsia"), config.get(font_key, "黑体"))

    def _generate_cover(self, doc, project_name, bidder_name, cover_data, config):
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        cover_lines = cover_data.get("cover_lines", [])
        if not cover_lines:
            cover_lines = [
                {"text": project_name or "投标文件", "style": "title"},
                {"text": "投 标 文 件", "style": "title"},
                {"text": "", "style": "normal"},
                {"text": f"投标人：{bidder_name or '（填写投标人名称）'}", "style": "normal"},
                {"text": "", "style": "normal"},
            ]

        for line_data in cover_lines:
            text = str(line_data.get("text", ""))
            style = line_data.get("style", "normal")

            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            if style == "title":
                run = para.add_run(text)
                run.font.size = Pt(config.get("title_size", 22))
                run.font.bold = True
                run.font.name = config.get("title_font", "方正小标宋简体")
                run.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("title_font", "方正小标宋简体"))
            elif style == "subtitle":
                run = para.add_run(text)
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.name = config.get("h1_font", "黑体")
                run.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("h1_font", "黑体"))
            else:
                run = para.add_run(text)
                run.font.size = Pt(14)
                run.font.name = config.get("body_font", "仿宋_GB2312")
                run.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("body_font", "仿宋_GB2312"))

        note_para = doc.add_paragraph()
        note_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note_run = note_para.add_run("此封面需加盖投标人公章")
        note_run.font.size = Pt(10)
        note_run.font.color.rgb = RGBColor(128, 128, 128)

        doc.add_page_break()

    def _add_chapters(self, doc, chapters, chapters_content, project_facts, config, level=1):
        from docx.shared import Pt
        from docx.oxml.ns import qn

        for idx, chapter in enumerate(chapters):
            if not isinstance(chapter, dict):
                continue

            title = chapter.get("title", "")
            ch_id = chapter.get("id", "")
            children = chapter.get("children", [])

            if level == 1:
                numbered_title = f"第{self._to_chinese_num(idx + 1)}章  {title}"
            else:
                numbered_title = f"{ch_id}  {title}" if ch_id else title

            heading_level = min(level, 4)
            heading = doc.add_heading(numbered_title, level=heading_level)
            heading.style.font.name = config.get(f"h{heading_level}_font", "黑体")
            heading.element.rPr.rFonts.set(qn("w:eastAsia"), config.get(f"h{heading_level}_font", "黑体"))

            content_key = ch_id
            content = chapters_content.get(content_key, "")
            if not content and not children:
                content = chapters_content.get(title, "")

            if content:
                content = self._fill_placeholders(content, project_facts)
                self._add_markdown_content(doc, content, config)

            if children:
                self._add_chapters(doc, children, chapters_content, project_facts, config, level + 1)

    def _add_markdown_content(self, doc, content, config):
        from docx.shared import Pt
        from docx.oxml.ns import qn

        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            if not line:
                i += 1
                continue

            if line.startswith("```"):
                i += 1
                code_lines = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                i += 1
                continue

            if line.startswith("|") and "|" in line[1:]:
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                self._add_table(doc, table_lines, config)
                continue

            if line.startswith("###"):
                text = line.lstrip("#").strip()
                h = doc.add_heading(text, level=3)
                h.style.font.name = config.get("h3_font", "黑体")
                h.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("h3_font", "黑体"))
            elif line.startswith("##"):
                text = line.lstrip("#").strip()
                h = doc.add_heading(text, level=2)
                h.style.font.name = config.get("h2_font", "黑体")
                h.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("h2_font", "黑体"))
            elif line.startswith("#"):
                text = line.lstrip("#").strip()
                h = doc.add_heading(text, level=1)
                h.style.font.name = config.get("h1_font", "黑体")
                h.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("h1_font", "黑体"))
            elif line.startswith("- ") or line.startswith("* "):
                text = self._clean_md(line[2:])
                para = doc.add_paragraph(text, style="List Bullet")
                self._set_body_font(para, config)
            elif re.match(r"^\d+\.\s", line):
                text = self._clean_md(re.sub(r"^\d+\.\s", "", line))
                para = doc.add_paragraph(text, style="List Number")
                self._set_body_font(para, config)
            else:
                text = self._clean_md(line)
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.name = config.get("english_font", "Times New Roman")
                run.font.size = Pt(config.get("body_size", 12))
                run.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("body_font", "仿宋_GB2312"))
                para.paragraph_format.first_line_indent = Pt(24)
                para.paragraph_format.line_spacing = Pt(config.get("line_spacing", 28))

            i += 1

    def _add_table(self, doc, table_lines, config):
        from docx.shared import Pt
        from docx.oxml.ns import qn

        if len(table_lines) < 2:
            return

        rows_data = []
        for line in table_lines:
            if re.match(r"^\|[\s\-:|]+\|$", line):
                continue
            cells = [c.strip() for c in line.split("|")[1:-1]]
            rows_data.append(cells)

        if not rows_data:
            return

        num_cols = max(len(row) for row in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.style = "Table Grid"

        for i, row_data in enumerate(rows_data):
            for j, cell_text in enumerate(row_data):
                if j < num_cols:
                    cell = table.cell(i, j)
                    cell.text = self._clean_md(cell_text)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(10.5)
                            run.font.name = config.get("english_font", "Times New Roman")
                            run.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("body_font", "仿宋_GB2312"))
                            if i == 0:
                                run.font.bold = True

    def _add_page_numbers(self, doc):
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        for section in doc.sections:
            footer = section.footer
            footer.is_linked_to_previous = False
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.alignment = 1  # CENTER

            run = para.add_run()
            fldChar1 = OxmlElement("w:fldChar")
            fldChar1.set(qn("w:fldCharType"), "begin")
            run._r.append(fldChar1)

            run2 = para.add_run()
            instrText = OxmlElement("w:instrText")
            instrText.set(qn("xml:space"), "preserve")
            instrText.text = " PAGE "
            run2._r.append(instrText)

            run3 = para.add_run()
            fldChar2 = OxmlElement("w:fldChar")
            fldChar2.set(qn("w:fldCharType"), "end")
            run3._r.append(fldChar2)

    def _set_body_font(self, para, config):
        from docx.shared import Pt
        from docx.oxml.ns import qn

        for run in para.runs:
            run.font.name = config.get("english_font", "Times New Roman")
            run.font.size = Pt(config.get("body_size", 12))
            run.element.rPr.rFonts.set(qn("w:eastAsia"), config.get("body_font", "仿宋_GB2312"))

    @staticmethod
    def _clean_md(text: str) -> str:
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        text = re.sub(r"`(.+?)`", r"\1", text)
        text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
        text = re.sub(r"~~(.+?)~~", r"\1", text)
        return text.strip()

    @staticmethod
    def _fill_placeholders(content: str, facts: dict) -> str:
        if not facts:
            return content

        replacements = {
            "本公司": facts.get("bidder_name", facts.get("company_name", "本公司")),
            "投标人": facts.get("bidder_name", facts.get("company_name", "投标人")),
            "XX公司": facts.get("bidder_name", facts.get("company_name", "")),
            "XX项目": facts.get("project_name", ""),
            "采购人": facts.get("purchaser_name", "采购人"),
            "招标人": facts.get("purchaser_name", "招标人"),
        }

        for placeholder, value in replacements.items():
            if value and placeholder != value:
                content = content.replace(f"[{placeholder}]", value)
                content = content.replace(f"\u3010{placeholder}\u3011", value)
                content = content.replace(f"${{{placeholder}}}", value)

        bidder = facts.get("bidder_name", facts.get("company_name", ""))
        if bidder:
            content = content.replace("XX科技有限公司", bidder)
            content = content.replace("XX有限公司", bidder)
        project_name = facts.get("project_name", "")
        if project_name:
            content = content.replace("XX项目", project_name)

        return content

    @staticmethod
    def _to_chinese_num(n: int) -> str:
        nums = "零一二三四五六七八九十"
        if n <= 10:
            return nums[n]
        elif n < 20:
            return f"十{nums[n - 10] if n > 10 else ''}"
        elif n < 100:
            tens = n // 10
            ones = n % 10
            return f"{nums[tens]}十{nums[ones] if ones else ''}"
        return str(n)
