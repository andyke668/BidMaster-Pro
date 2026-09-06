from __future__ import annotations

from core.doc_engine.parsers.base import ParsedDocument

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"


class DocxParser:
    def parse(self, file_path: str) -> ParsedDocument:
        import os

        if not os.path.exists(file_path):
            raise ValueError(f"文件不存在: {file_path}（请重新上传）")

        with open(file_path, "rb") as f:
            head = f.read(8)

        if head[:4] == _ZIP_MAGIC:
            return self._parse_docx(file_path)

        if head == _OLE2_MAGIC:
            # 老式 Word 97-2003 二进制 .doc（或 WPS 保存的变体）：python-docx 打不开
            return self._parse_legacy_doc(file_path)

        if head[:4] == b"%PDF":
            # 扩展名与内容不符：实际是 PDF，交给 PdfParser
            from core.doc_engine.parsers.pdf_parser import PdfParser
            return PdfParser().parse(file_path)

        raise ValueError(
            "无法识别的文件内容（不是 docx/doc/pdf）。请在 Word/WPS 中另存为 .docx 后重新上传"
        )

    def _parse_docx(self, file_path: str) -> ParsedDocument:
        from docx import Document

        doc = Document(file_path)
        text_parts = []
        tables = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                style_name = para.style.name if para.style else ""
                text_parts.append(text)

        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            tables.append({
                "index": i,
                "headers": rows[0] if rows else [],
                "rows": rows[1:] if len(rows) > 1 else [],
            })

        return ParsedDocument(
            text="\n".join(text_parts),
            tables=tables,
            images=[],
            metadata={
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "parser": "python-docx",
            },
        )

    def _parse_legacy_doc(self, file_path: str) -> ParsedDocument:
        """老式 Word 97-2003 二进制 .doc：antiword 优先，失败回退 catdoc。

        antiword 2005 年后未更新，对 WPS Office 写出的 .doc（FIB 头非标准）会直接
        拒绝；catdoc 兼容性更好，实测可正确提取中文。两者都只输出纯文本，
        表格无法结构化还原（tables 为空）。
        """
        import shutil
        import subprocess

        errors: list[str] = []
        text = ""
        used = ""
        for tool, args in (
            ("antiword", ["antiword", "-m", "UTF-8.txt", file_path]),
            ("catdoc", ["catdoc", "-d", "utf-8", file_path]),
        ):
            if not shutil.which(tool):
                errors.append(f"{tool} 未安装")
                continue
            try:
                proc = subprocess.run(args, capture_output=True, timeout=120)
            except subprocess.TimeoutExpired:
                errors.append(f"{tool} 解析超时")
                continue
            if proc.returncode != 0:
                err = proc.stderr.decode("utf-8", errors="replace").strip()[:120]
                errors.append(f"{tool}: {err or '无法读取'}")
                continue
            out = proc.stdout.decode("utf-8", errors="replace")
            if out.strip():
                text = out
                used = tool
                break
            errors.append(f"{tool}: 未提取到文本")

        if not text.strip():
            raise ValueError(
                "旧版 .doc 解析失败（文件可能加密、损坏或为纯扫描件）："
                + "；".join(errors)
                + "。请在 Word/WPS 中另存为 .docx 后重新上传"
            )

        lines = [ln.rstrip() for ln in text.splitlines()]
        # catdoc 对快速保存的文件会输出一行英文警告，去掉；同时过滤空行
        lines = [ln for ln in lines if ln.strip() and not ln.startswith("[This was fast-saved")]
        text = "\n".join(lines)

        return ParsedDocument(
            text=text,
            tables=[],
            images=[],
            metadata={
                "parser": used,
                "legacy_doc": True,
                "char_count": len(text),
            },
        )